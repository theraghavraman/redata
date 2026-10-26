#!/usr/bin/env python3
"""ReData local engine.

Runs on your own computer, serves ReData at http://127.0.0.1:8765 and does all
heavy data work (import, profiling, sorting, aggregation, statistics, SQL,
cleaning, export) in DuckDB instead of the browser tab. The browser only ever
receives small results: one page of rows, one chart's aggregates, one table of
statistics. That is what lets ReData work with crores of rows.

Nothing leaves your machine: the server only listens on the loopback address,
checks the Host header, and every API call needs a random token that is only
handed to the ReData page this engine serves.
"""

import argparse
import csv
import datetime
import decimal
import json
import math
import os
import re
import secrets
import shutil
import sys
import threading
import time
import uuid
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, unquote, urlparse

try:
    import duckdb
except ImportError:  # pragma: no cover - friendly message for first run
    sys.exit(
        "DuckDB is not installed. Run the start script (start-engine.bat / "
        "start-engine.sh) or: python -m pip install -r engine/requirements.txt"
    )

VERSION = "1.1.0"
ENGINE_DIR = os.path.dirname(os.path.abspath(__file__))
APP_ROOT = os.path.dirname(ENGINE_DIR)
TABLE = "data"
PAGE_MAX = 500
PREVIEW_ROWS = 20
SQL_UI_ROWS = 1000
EXACT_DISTINCT_LIMIT = 2_000_000
CHUNK = 1024 * 1024

NUMERIC_PREFIXES = (
    "TINYINT", "SMALLINT", "INTEGER", "BIGINT", "HUGEINT", "UTINYINT",
    "USMALLINT", "UINTEGER", "UBIGINT", "UHUGEINT", "FLOAT", "DOUBLE",
    "DECIMAL", "REAL",
)
DATE_PREFIXES = ("DATE", "TIMESTAMP", "TIME")

# Table functions that read from the filesystem or network. The SQL lab is a
# read-only *query* surface over the loaded dataset, not a file browser, so
# these are rejected even inside an otherwise-valid SELECT (incl. CTEs/subqueries).
FILE_FUNCTIONS = (
    "read_csv", "read_csv_auto", "read_json", "read_json_auto", "read_ndjson",
    "read_ndjson_auto", "read_json_objects", "read_parquet", "parquet_scan",
    "parquet_metadata", "parquet_schema", "parquet_file_metadata",
    "parquet_kv_metadata", "read_text", "read_blob", "glob", "sniff_csv",
    "read_xlsx", "delta_scan", "iceberg_scan",
)
_FILE_FUNC_RE = re.compile(r"\b(" + "|".join(FILE_FUNCTIONS) + r")\s*\(", re.IGNORECASE)
STATIC_TYPES = {
    ".html": "text/html; charset=utf-8",
    ".svg": "image/svg+xml",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".ico": "image/x-icon",
    ".css": "text/css; charset=utf-8",
    ".js": "text/javascript; charset=utf-8",
    ".woff2": "font/woff2",
}


class UserError(Exception):
    """An error whose message is safe and useful to show in the UI."""


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def qi(name):
    """Quote an SQL identifier."""
    return '"' + str(name).replace('"', '""') + '"'


def is_numeric_type(t):
    return str(t).upper().startswith(NUMERIC_PREFIXES)


def is_date_type(t):
    return str(t).upper().startswith(DATE_PREFIXES)


def is_text_type(t):
    return str(t).upper() == "VARCHAR"


def is_nested_type(t):
    u = str(t).upper()
    return u.startswith(("STRUCT", "MAP", "UNION")) or u.endswith("]")


def plain(v):
    """Convert a DuckDB value into something json.dumps accepts."""
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        return v if math.isfinite(v) else None
    if isinstance(v, decimal.Decimal):
        f = float(v)
        return f if math.isfinite(f) else None
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, datetime.timedelta):
        return str(v)
    if isinstance(v, (bytes, bytearray)):
        return v.hex()
    if isinstance(v, (list, tuple)):
        return [plain(x) for x in v]
    if isinstance(v, dict):
        return {str(k): plain(x) for k, x in v.items()}
    return str(v)


def unique_headers(raw, count):
    """Same rules as the browser: blank -> column_N, duplicates -> name_2 (case-insensitive)."""
    seen, out = set(), []
    for i in range(count):
        v = raw[i] if raw and i < len(raw) and raw[i] is not None else ""
        h = str(v).strip() or "column_%d" % (i + 1)
        base, n = h, 2
        while h.lower() in seen:
            h = "%s_%d" % (base, n)
            n += 1
        seen.add(h.lower())
        out.append(h)
    return out


def safe_filename(name):
    name = os.path.basename(str(name or "dataset"))
    name = re.sub(r"[^A-Za-z0-9._ -]+", "_", name).strip(" .") or "dataset"
    return name[:120]


def file_ext(name):
    return os.path.splitext(str(name))[1].lower().lstrip(".")


# --------------------------------------------------------------------------
# formula -> SQL (same grammar as the browser formula builder)
# --------------------------------------------------------------------------

_TOKEN_RE = re.compile(
    r"\s*(?:(\d+(?:\.\d*)?|\.\d+)([eE][+-]?\d+)?|(@\d+)|([()+\-*/%]))"
)


def formula_to_sql(expression, columns):
    refs = []

    def sub(m):
        name = m.group(1).strip()
        if name not in columns:
            raise UserError("Unknown column reference: " + name)
        refs.append(name)
        return "@%d" % (len(refs) - 1)

    src = re.sub(r"\[([^\]]+)\]", sub, expression)
    toks, pos = [], 0
    while pos < len(src):
        if src[pos:].strip() == "":
            break
        m = _TOKEN_RE.match(src, pos)
        if not m or m.end() == pos:
            raise UserError("Invalid formula near: " + src[pos:pos + 12])
        pos = m.end()
        if m.group(1) is not None:
            toks.append(("num", float(m.group(1) + (m.group(2) or ""))))
        elif m.group(3):
            toks.append(("ref", refs[int(m.group(3)[1:])]))
        else:
            toks.append((m.group(4), None))
    i = 0

    def peek():
        return toks[i][0] if i < len(toks) else None

    def atom():
        nonlocal i
        if i >= len(toks):
            raise UserError("Unexpected end of formula.")
        t, v = toks[i]
        i += 1
        if t == "num":
            return "CAST(%r AS DOUBLE)" % v
        if t == "ref":
            return "TRY_CAST(%s AS DOUBLE)" % qi(v)
        if t in ("-", "+"):
            inner = atom()
            return "(-(%s))" % inner if t == "-" else inner
        if t == "(":
            v2 = expr()
            if peek() != ")":
                raise UserError("Missing closing parenthesis.")
            i += 1
            return "(%s)" % v2
        raise UserError("Expected a number, column reference or parenthesis.")

    def term():
        nonlocal i
        v = atom()
        while peek() in ("*", "/", "%"):
            op = toks[i][0]
            i += 1
            n = atom()
            if op == "*":
                v = "(%s * %s)" % (v, n)
            elif op == "/":
                v = "(%s / NULLIF(%s, 0))" % (v, n)
            else:
                v = "fmod(%s, NULLIF(%s, 0))" % (v, n)
        return v

    def expr():
        nonlocal i
        v = term()
        while peek() in ("+", "-"):
            op = toks[i][0]
            i += 1
            v = "(%s %s %s)" % (v, op, term())
        return v

    if not toks:
        raise UserError("Enter an expression.")
    out = expr()
    if i != len(toks):
        raise UserError("Unexpected token in formula.")
    return out


# --------------------------------------------------------------------------
# the engine
# --------------------------------------------------------------------------

class Engine:
    def __init__(self, workspace, memory_limit=None, threads=None):
        self.workspace = workspace
        self.stage_dir = os.path.join(workspace, "stage")
        self.export_dir = os.path.join(workspace, "exports")
        self.temp_dir = os.path.join(workspace, "tmp")
        for d in (workspace, self.stage_dir, self.export_dir, self.temp_dir):
            os.makedirs(d, exist_ok=True)
        # leftovers from an earlier run are never needed again
        for d in (self.stage_dir, self.export_dir, self.temp_dir):
            for entry in os.listdir(d):
                shutil.rmtree(os.path.join(d, entry), ignore_errors=True)
        self.state_path = os.path.join(workspace, "state.json")
        self.con = duckdb.connect(os.path.join(workspace, "redata.duckdb"))
        self.con.execute("SET temp_directory = ?", [self.temp_dir])
        self.con.execute("SET preserve_insertion_order = true")
        if memory_limit:
            self.con.execute("SET memory_limit = ?", [memory_limit])
        if threads:
            self.con.execute("SET threads = %d" % int(threads))
        self.lock = threading.RLock()
        self.stages = {}
        self.exports = {}
        self.meta_cache = None
        self.data_version = 0
        self.pipeline = None
        self.state = {"name": None, "linked_stream": None, "updated_at": None}
        try:
            with open(self.state_path, "r", encoding="utf-8") as fh:
                self.state.update(json.load(fh))
        except (OSError, ValueError):
            pass
        if not self.has_data():
            self.state["name"] = None

    def start_pipeline(self):
        from pipeline import Pipeline
        self.pipeline = Pipeline(self)

    # ---- basics ----------------------------------------------------------
    def q(self, sql, params=None):
        return self.con.execute(sql, params or []).fetchall()

    def has_data(self):
        return bool(self.q(
            "SELECT 1 FROM information_schema.tables WHERE table_schema='main' AND table_name=?",
            [TABLE]))

    def require_data(self):
        if not self.has_data():
            raise UserError("No dataset is loaded in the local engine yet. Import a file first.")

    def columns(self):
        return [(r[0], r[1]) for r in self.q("DESCRIBE " + qi(TABLE))]

    def col_names(self):
        return [c for c, _ in self.columns()]

    def col_type(self, name):
        for c, t in self.columns():
            if c == name:
                return t
        raise UserError("Unknown column: " + str(name))

    def check_col(self, name):
        self.col_type(name)
        return qi(name)

    def row_count(self):
        return self.q("SELECT count(*) FROM " + qi(TABLE))[0][0]

    def save_state(self):
        try:
            with open(self.state_path, "w", encoding="utf-8") as fh:
                json.dump(self.state, fh)
        except OSError:
            pass

    def changed(self):
        self.meta_cache = None
        self.data_version += 1
        self.state["updated_at"] = datetime.datetime.now().astimezone().isoformat(timespec="seconds")
        self.save_state()

    def rebuild(self, select_list, where="", extra_from=""):
        """Rewrite the table, keeping original row order and contiguous rowids."""
        self.q(
            "CREATE OR REPLACE TABLE %s AS SELECT %s FROM (SELECT *, rowid AS __redata_rid FROM %s) %s %s ORDER BY __redata_rid"
            % (qi(TABLE), select_list, qi(TABLE), extra_from, where)
        )
        self.changed()

    # ---- status / meta -----------------------------------------------------
    def status(self):
        running = 0
        if self.pipeline:
            running = sum(1 for v in self.pipeline.status.values() if v == "running")
        return {"version": VERSION, "dataset": self.state.get("name") if self.has_data() else None,
                "dataVersion": self.data_version, "updatedAt": self.state.get("updated_at"),
                "linkedStream": self.state.get("linked_stream"), "syncsRunning": running}

    # ---- staging tables as the active dataset ---------------------------
    def open_staged(self, table, link=False):
        from pipeline import SCHEMA
        exists = self.q("SELECT 1 FROM duckdb_tables() WHERE schema_name = ? AND table_name = ?", [SCHEMA, table])
        if not exists:
            raise UserError("No staging table called %s." % table)
        self.q("CREATE OR REPLACE TABLE %s AS SELECT * FROM %s.%s" % (qi(TABLE), SCHEMA, qi(table)))
        stream = None
        if link and self.pipeline:
            stream = next((x["id"] for x in self.pipeline.cfg["streams"] if x["target"] == table), None)
        self.state["name"] = "staging.%s" % table
        self.state["linked_stream"] = stream
        self.changed()
        return self.meta()

    def refresh_linked(self):
        from pipeline import SCHEMA
        sid = self.state.get("linked_stream")
        if not sid or not self.pipeline:
            return
        target = next((x["target"] for x in self.pipeline.cfg["streams"] if x["id"] == sid), None)
        if not target:
            return
        with self.lock:
            self.q("CREATE OR REPLACE TABLE %s AS SELECT * FROM %s.%s" % (qi(TABLE), SCHEMA, qi(target)))
            self.changed()
            self.meta()

    def excel_file_to_csv(self, path):
        st = {"path": path, "ext": file_ext(path)}
        return self._excel_to_csv(st, 0)

    def meta(self):
        if self.meta_cache is not None:
            return self.meta_cache
        self.require_data()
        cols = self.columns()
        n = self.row_count()
        exact = n <= EXACT_DISTINCT_LIMIT
        parts, layout = [], []
        for c, t in cols:
            id_ = qi(c)
            parts.append("count(%s)" % id_)
            parts.append(("count(DISTINCT %s)" if exact else "approx_count_distinct(%s)") % id_)
            if is_numeric_type(t):
                parts.append(
                    "avg(%s), min(%s), max(%s), stddev_samp(%s), quantile_cont(%s, [0.25, 0.5, 0.75])"
                    % (id_, id_, id_, id_, id_))
                layout.append((c, t, "num"))
            else:
                parts.append("approx_top_k(%s, 1)" % id_)
                layout.append((c, t, "cat"))
        row = self.q("SELECT %s FROM %s" % (", ".join(parts), qi(TABLE)))[0] if parts else []
        out, k = [], 0
        for c, t, kind in layout:
            nonnull, distinct = row[k], row[k + 1]
            k += 2
            info = {
                "name": c, "type": t, "numeric": kind == "num",
                "date": is_date_type(t) or bool(re.search("date|time", c, re.I)),
                "nonEmpty": nonnull, "missing": n - nonnull,
                "distinct": distinct, "distinctExact": exact,
            }
            if kind == "num":
                avg, mn, mx, sd, qs = row[k:k + 5]
                k += 5
                qs = qs or [None, None, None]
                info.update({
                    "mean": plain(avg), "min": plain(mn), "max": plain(mx),
                    "sd": plain(sd) if nonnull >= 2 else 0,
                    "q1": plain(qs[0]), "median": plain(qs[1]), "q3": plain(qs[2]),
                })
            else:
                top = row[k]
                k += 1
                info["top"] = plain(top[0]) if top else None
                info["topApprox"] = (distinct or 0) > 1000
            out.append(info)
        # approximate distinct counts are replaced by exact ones where cheap
        if not exact:
            small = [i for i in out if (i["distinct"] or 0) <= 100000]
            if small:
                res = self.q("SELECT %s FROM %s" % (
                    ", ".join("count(DISTINCT %s)" % qi(i["name"]) for i in small), qi(TABLE)))[0]
                for i, v in zip(small, res):
                    i["distinct"], i["distinctExact"] = v, True
        # exact most-frequent value wherever the column has few enough values
        modal = [i for i in out if not i["numeric"] and i["distinctExact"] and (i["distinct"] or 0) <= 100000]
        if modal:
            res = self.q("SELECT %s FROM %s" % (
                ", ".join("mode(%s)" % qi(i["name"]) for i in modal), qi(TABLE)))[0]
            for i, v in zip(modal, res):
                i["top"], i["topApprox"] = plain(v), False
        for i in out:
            if not i["numeric"] and i not in modal:
                i["topApprox"] = True
        # Tukey IQR outlier counts need the quartiles first
        nums = [i for i in out if i["numeric"] and i["q1"] is not None]
        if nums:
            exprs = []
            for i in nums:
                iqr = i["q3"] - i["q1"]
                lo, hi = i["q1"] - 1.5 * iqr, i["q3"] + 1.5 * iqr
                exprs.append("count(*) FILTER (WHERE %s < %r OR %s > %r)" % (qi(i["name"]), lo, qi(i["name"]), hi))
            res = self.q("SELECT %s FROM %s" % (", ".join(exprs), qi(TABLE)))[0]
            for i, v in zip(nums, res):
                i["outliers"] = v
        for i in out:
            if i["numeric"]:
                i.setdefault("outliers", 0)
        dups = 0
        if cols:
            distinct_rows = self.q("SELECT count(*) FROM (SELECT DISTINCT * FROM %s)" % qi(TABLE))[0][0]
            dups = n - distinct_rows
        self.meta_cache = {
            "name": self.state.get("name") or "dataset",
            "version": self.data_version,
            "updatedAt": self.state.get("updated_at"),
            "linkedStream": self.state.get("linked_stream"),
            "rows": n,
            "columns": out,
            "missingTotal": sum(i["missing"] for i in out),
            "duplicates": dups,
        }
        return self.meta_cache

    # ---- staging ------------------------------------------------------------
    def stage_upload(self, filename, stream, length):
        filename = safe_filename(filename)
        ext = file_ext(filename)
        self._check_ext(ext)
        sid = uuid.uuid4().hex
        folder = os.path.join(self.stage_dir, sid)
        os.makedirs(folder)
        path = os.path.join(folder, filename)
        remaining = length
        with open(path, "wb") as fh:
            while remaining > 0:
                chunk = stream.read(min(CHUNK, remaining))
                if not chunk:
                    break
                fh.write(chunk)
                remaining -= len(chunk)
        if remaining > 0:
            shutil.rmtree(folder, ignore_errors=True)
            raise UserError("The upload was interrupted before the whole file arrived.")
        with self.lock:
            return self._register_stage(sid, path, filename, ext, copied=True, folder=folder)

    def stage_path(self, raw_path):
        path = os.path.abspath(os.path.expanduser(str(raw_path or "").strip().strip('"').strip("'")))
        if not raw_path or not os.path.isfile(path):
            raise UserError("No file was found at that path on this computer.")
        ext = file_ext(path)
        self._check_ext(ext)
        return self._register_stage(uuid.uuid4().hex, path, os.path.basename(path), ext, copied=False, folder=None)

    def _check_ext(self, ext):
        if ext not in ("csv", "tsv", "txt", "json", "ndjson", "jsonl", "xlsx", "xlsm", "xls", "parquet"):
            raise UserError("Unsupported file type: ." + (ext or "?") +
                            " (use CSV, TSV, TXT, JSON, NDJSON, Parquet, XLSX or XLS)")

    def _register_stage(self, sid, path, name, ext, copied, folder):
        kind = {"csv": "delimited", "tsv": "delimited", "txt": "delimited",
                "json": "json", "ndjson": "json", "jsonl": "json",
                "xlsx": "excel", "xlsm": "excel", "xls": "excel",
                "parquet": "parquet"}[ext]
        st = {"id": sid, "path": path, "name": name, "ext": ext, "kind": kind,
              "copied": copied, "folder": folder, "size": os.path.getsize(path)}
        self.stages[sid] = st
        resp = {"stage": sid, "name": name, "kind": kind, "size": st["size"], "sheets": [], "preview": None}
        if kind == "delimited":
            resp["preview"] = self._csv_preview(st)
        elif kind == "excel":
            sheets = self._excel_sheets(st)
            resp["sheets"] = sheets
            resp["preview"] = self.excel_preview(sid, 0)
        return resp

    def _stage(self, sid):
        st = self.stages.get(sid)
        if not st or not os.path.isfile(st["path"]):
            raise UserError("The selected file is no longer available. Choose it again.")
        return st

    def drop_stage(self, sid):
        st = self.stages.pop(sid, None)
        if st and st["copied"] and st["folder"]:
            shutil.rmtree(st["folder"], ignore_errors=True)

    def _csv_source(self, st, header, sample_size=None, all_varchar=False):
        opts = ["header=%s" % ("true" if header else "false"), "null_padding=true"]
        if st["ext"] == "tsv":
            opts.append("delim='\\t'")
        if st.get("parallel") is False:
            # DuckDB's parallel CSV scanner cannot combine null_padding with
            # quoted newlines. Set this only after preview detects that specific
            # compatibility error, preserving parallel imports for ordinary files.
            opts.append("parallel=false")
        if sample_size is not None:
            opts.append("sample_size=%d" % sample_size)
        if all_varchar:
            opts.append("all_varchar=true")
        return "read_csv(%s, %s)" % (self._lit(st["path"]), ", ".join(opts))

    def _delimited_headers(self, st):
        """Read original header cells before DuckDB normalizes duplicate names."""
        try:
            with open(st["path"], "r", encoding="utf-8-sig", newline="") as stream:
                sample = stream.read(65536)
                if st["ext"] == "tsv":
                    delimiter = "\t"
                elif st["ext"] == "csv":
                    delimiter = ","
                else:
                    try:
                        delimiter = csv.Sniffer().sniff(
                            sample[:8192], delimiters=",\t;|"
                        ).delimiter
                    except csv.Error:
                        delimiter = ","
                stream.seek(0)
                return next(csv.reader(stream, delimiter=delimiter), [])
        except (OSError, UnicodeError, csv.Error):
            # Preserve DuckDB's normal fallback for unusual encodings/dialects.
            return None

    @staticmethod
    def _lit(s):
        return "'" + str(s).replace("'", "''") + "'"

    def _csv_preview(self, st):
        sql = "SELECT * FROM %s LIMIT %d" % (
            self._csv_source(st, False, all_varchar=True), PREVIEW_ROWS
        )
        try:
            rows = self.q(sql)
        except duckdb.Error as exc:
            message = str(exc).lower()
            if "null_padding" not in message or "quoted new lines" not in message:
                raise
            st["parallel"] = False
            sql = "SELECT * FROM %s LIMIT %d" % (
                self._csv_source(st, False, all_varchar=True), PREVIEW_ROWS
            )
            rows = self.q(sql)
        return [[plain(v) for v in r] for r in rows]

    # ---- excel ----------------------------------------------------------------
    def _excel_sheets(self, st):
        if st["ext"] == "xls":
            xlrd = self._xlrd()
            book = xlrd.open_workbook(st["path"], on_demand=True)
            return book.sheet_names()
        openpyxl = self._openpyxl()
        wb = openpyxl.load_workbook(st["path"], read_only=True, data_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()

    def _excel_rows(self, st, sheet_index):
        if st["ext"] == "xls":
            xlrd = self._xlrd()
            book = xlrd.open_workbook(st["path"], on_demand=True)
            sh = book.sheet_by_index(sheet_index)
            for r in range(sh.nrows):
                yield sh.row_values(r)
            return
        openpyxl = self._openpyxl()
        wb = openpyxl.load_workbook(st["path"], read_only=True, data_only=True)
        try:
            ws = wb.worksheets[sheet_index]
            for row in ws.iter_rows(values_only=True):
                yield row
        finally:
            wb.close()

    def excel_preview(self, sid, sheet_index):
        st = self._stage(sid)
        out = []
        for row in self._excel_rows(st, int(sheet_index or 0)):
            vals = [plain(v) for v in row]
            if any(v is not None and str(v).strip() != "" for v in vals):
                out.append(vals)
            if len(out) >= PREVIEW_ROWS:
                break
        return out

    @staticmethod
    def _openpyxl():
        try:
            import openpyxl  # noqa: F401
            return openpyxl
        except ImportError:
            raise UserError("Excel import needs the openpyxl package: python -m pip install openpyxl")

    @staticmethod
    def _xlrd():
        try:
            import xlrd  # noqa: F401
            return xlrd
        except ImportError:
            raise UserError("Old .xls files need the xlrd package: python -m pip install xlrd "
                            "(or save the sheet as .xlsx or .csv)")

    def _excel_to_csv(self, st, sheet_index):
        out_path = os.path.join(self.temp_dir, uuid.uuid4().hex + ".csv")
        with open(out_path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            for row in self._excel_rows(st, sheet_index):
                w.writerow(["" if v is None else plain(v) for v in row])
        return out_path

    # ---- loading -----------------------------------------------------------------
    def load(self, sid, header=True, names=None, sheet=0):
        st = self._stage(sid)
        tmp_csv = None
        try:
            if st["kind"] == "delimited":
                source = self._pick_csv_source(st, header)
                display = st["name"]
                if header and not names:
                    names = self._delimited_headers(st)
            elif st["kind"] == "excel":
                tmp_csv = self._excel_to_csv(st, int(sheet or 0))
                fake = dict(st, path=tmp_csv, ext="csv")
                source = self._pick_csv_source(fake, header)
                sheets = self._excel_sheets(st)
                display = st["name"] + " · " + sheets[int(sheet or 0)]
            elif st["kind"] == "json":
                source = self._json_source(st)
                display = st["name"]
                names = None
            else:
                source = "read_parquet(%s)" % self._lit(st["path"])
                display = st["name"]
                names = None
            self._create_from(source, names)
        finally:
            if tmp_csv:
                try:
                    os.remove(tmp_csv)
                except OSError:
                    pass
        self.drop_stage(sid)
        self.state["name"] = display
        self.state["linked_stream"] = None
        self.save_state()
        self.changed()
        return self.meta()

    def _pick_csv_source(self, st, header):
        """DuckDB guesses column types from a sample. If a value later in the file
        does not fit the guess, the load is retried with a full-file guess, then as text."""
        return [self._csv_source(st, header),
                self._csv_source(st, header, sample_size=-1),
                self._csv_source(st, header, all_varchar=True)]

    def _json_source(self, st):
        src = "read_json_auto(%s, format='auto', maximum_object_size=1073741824)" % self._lit(st["path"])
        try:
            desc = self.q("DESCRIBE SELECT * FROM %s" % src)
            # {"rows": [ {...}, ... ]} -> unwrap the first array of objects
            if len(desc) >= 1:
                lists = [(c, t) for c, t, *_ in desc if str(t).upper().startswith("STRUCT") and str(t).endswith("[]")]
                if lists and self.q("SELECT count(*) FROM %s" % src)[0][0] == 1:
                    c = lists[0][0]
                    return "(SELECT u.* FROM (SELECT unnest(%s) AS u FROM %s))" % (qi(c), src)
        except duckdb.Error as exc:
            raise UserError("Could not read the JSON file: %s" % str(exc).split("\n")[0])
        return src

    def _create_from(self, sources, names):
        if isinstance(sources, str):
            sources = [sources]
        last = None
        for i, source in enumerate(sources):
            try:
                return self._create_one(source, names)
            except UserError:
                raise
            except duckdb.Error as e:
                last = e
                self.q("DROP TABLE IF EXISTS __redata_load")
                if i < len(sources) - 1:
                    sys.stdout.write("  type guess did not fit every row, retrying (%s)\n" % str(e).split("\n")[0][:120])
        raise UserError("Could not read the file: %s" % str(last).split("\n")[0])

    def _create_one(self, source, names):
        desc = self.q("DESCRIBE SELECT * FROM %s" % source)
        if not desc:
            raise UserError("No columns were found in the file.")
        src_cols = [(r[0], r[1]) for r in desc]
        final = unique_headers(names if names else [c for c, _ in src_cols], len(src_cols))
        select, not_all_null = [], []
        for (c, t), alias in zip(src_cols, final):
            id_ = qi(c)
            if is_text_type(t):
                e = "NULLIF(trim(%s), '')" % id_
            elif is_nested_type(t):
                e = "CAST(to_json(%s) AS VARCHAR)" % id_
            else:
                e = id_
            select.append("%s AS %s" % (e, qi(alias)))
            not_all_null.append("%s IS NOT NULL" % e)
        sql = "CREATE OR REPLACE TABLE __redata_load AS SELECT %s FROM %s WHERE %s" % (
            ", ".join(select), source, " OR ".join(not_all_null))
        self.q(sql)
        n = self.q("SELECT count(*) FROM __redata_load")[0][0]
        if not n:
            self.q("DROP TABLE __redata_load")
            raise UserError("No data rows were found after skipping blank rows.")
        self.q("DROP TABLE IF EXISTS %s" % qi(TABLE))
        self.q("ALTER TABLE __redata_load RENAME TO %s" % qi(TABLE))

    # ---- reading -----------------------------------------------------------------
    def rows(self, offset=0, limit=150, sort=None, direction=1, column=None, query=""):
        self.require_data()
        offset = max(0, int(offset or 0))
        limit = max(1, min(PAGE_MAX, int(limit or 150)))
        where, params = "", []
        if column and query:
            where = "WHERE contains(lower(CAST(%s AS VARCHAR)), lower(?))" % self.check_col(column)
            params.append(str(query))
        n = self.row_count()
        if not where and not sort:
            # rowids are kept contiguous, so a page is a direct range read
            res = self.con.execute(
                "SELECT * FROM %s WHERE rowid >= ? AND rowid < ? ORDER BY rowid" % qi(TABLE),
                [offset, offset + limit])
            total = n
        else:
            order = "rowid"
            if sort:
                order = "%s %s NULLS LAST, rowid" % (self.check_col(sort), "DESC" if int(direction or 1) < 0 else "ASC")
            res = self.con.execute(
                "SELECT * FROM %s %s ORDER BY %s LIMIT %d OFFSET %d" % (qi(TABLE), where, order, limit, offset),
                params)
            total = n if not where else None
        cols = [d[0] for d in res.description]
        data = [dict(zip(cols, [plain(v) for v in r])) for r in res.fetchall()]
        if total is None:
            total = self.q("SELECT count(*) FROM %s %s" % (qi(TABLE), where), params)[0][0]
        return {"columns": cols, "rows": data, "total": total, "offset": offset}

    @staticmethod
    def _agg_expr(agg, val):
        return {"sum": "sum(%s)", "mean": "avg(%s)", "median": "median(%s)"}.get(agg, "count(*)") % ((val,) if agg in ("sum", "mean", "median") else ())

    def aggregate(self, g, y=None, agg="sum", order="desc", limit=20, with_groups=False):
        self.require_data()
        gq = self.check_col(g)
        agg = agg if agg in ("sum", "mean", "median", "count") else "sum"
        val = "TRY_CAST(%s AS DOUBLE)" % self.check_col(y) if (y and agg != "count") else "NULL"
        value_expr = self._agg_expr(agg, "__v")
        having = "" if agg == "count" else "HAVING count(__v) > 0"
        limit = max(1, min(5000, int(limit or 20)))
        sql = ("SELECT COALESCE(CAST(%s AS VARCHAR), '(blank)') AS label, count(*) AS rows, %s AS n, %s AS value "
               "FROM (SELECT %s, %s AS __v FROM %s) GROUP BY 1 %s ORDER BY value %s NULLS LAST, label LIMIT %d") % (
            gq, "count(*)" if agg == "count" else "count(__v)", value_expr, gq, val, qi(TABLE), having,
            "ASC" if order == "asc" else "DESC", limit)
        res = self.q(sql)
        out = {"rows": [{"label": r[0], "rows": r[1], "n": r[2], "value": plain(r[3])} for r in res]}
        if with_groups:
            out["groups"] = self.q(
                "SELECT count(*) FROM (SELECT 1 FROM (SELECT %s, %s AS __v FROM %s) GROUP BY %s %s)"
                % (gq, val, qi(TABLE), gq, having))[0][0]
        return out

    def _aggregate_export_sql(self, g, y, agg):
        gq = self.check_col(g)
        agg = agg if agg in ("sum", "mean", "median", "count") else "sum"
        val = "TRY_CAST(%s AS DOUBLE)" % self.check_col(y) if (y and agg != "count") else "NULL"
        having = "" if agg == "count" else "HAVING count(__v) > 0"
        return ("SELECT COALESCE(CAST(%s AS VARCHAR), '(blank)') AS \"Group\", count(*) AS \"Rows\", %s AS \"Valid\", %s AS \"Value\" "
                "FROM (SELECT %s, %s AS __v FROM %s) GROUP BY 1 %s ORDER BY \"Value\" DESC NULLS LAST") % (
            gq, "count(*)" if agg == "count" else "count(__v)", self._agg_expr(agg, "__v"), gq, val, qi(TABLE), having)

    def histogram(self, col, bins=10):
        self.require_data()
        c = "TRY_CAST(%s AS DOUBLE)" % self.check_col(col)
        bins = max(1, min(100, int(bins or 10)))
        lo, hi, n = self.q("SELECT min(v), max(v), count(v) FROM (SELECT %s AS v FROM %s) WHERE isfinite(v)" % (c, qi(TABLE)))[0]
        if not n:
            return {"lo": None, "hi": None, "counts": []}
        width = (hi - lo) / bins or 1
        res = dict(self.q(
            "SELECT least(%d, CAST(floor((v - ?) / ?) AS BIGINT)) AS b, count(*) FROM (SELECT %s AS v FROM %s) "
            "WHERE isfinite(v) GROUP BY 1" % (bins - 1, c, qi(TABLE)), [lo, width]))
        return {"lo": lo, "hi": hi, "counts": [res.get(i, 0) for i in range(bins)]}

    def box(self, col):
        self.require_data()
        c = "TRY_CAST(%s AS DOUBLE)" % self.check_col(col)
        lo, hi, qs, n = self.q(
            "SELECT min(v), max(v), quantile_cont(v, [0.25, 0.5, 0.75]), count(v) FROM (SELECT %s AS v FROM %s) WHERE isfinite(v)"
            % (c, qi(TABLE)))[0]
        if not n:
            return {"n": 0}
        q1, med, q3 = qs
        iqr = q3 - q1
        wl, wh = max(lo, q1 - 1.5 * iqr), min(hi, q3 + 1.5 * iqr)
        outs = [r[0] for r in self.q(
            "SELECT v FROM (SELECT %s AS v FROM %s) WHERE v < ? OR v > ? ORDER BY v LIMIT 250" % (c, qi(TABLE)), [wl, wh])]
        return {"n": n, "lo": lo, "hi": hi, "q1": q1, "med": med, "q3": q3, "outliers": outs}

    def scatter(self, x, y, limit=5000):
        self.require_data()
        xq, yq = self.check_col(x), self.check_col(y)
        limit = max(10, min(20000, int(limit or 5000)))
        res = self.q(
            "SELECT x, y FROM (SELECT TRY_CAST(%s AS DOUBLE) AS x, TRY_CAST(%s AS DOUBLE) AS y FROM %s) "
            "WHERE isfinite(x) AND isfinite(y) USING SAMPLE %d ROWS" % (xq, yq, qi(TABLE), limit))
        total = self.q(
            "SELECT count(*) FROM (SELECT TRY_CAST(%s AS DOUBLE) AS x, TRY_CAST(%s AS DOUBLE) AS y FROM %s) WHERE isfinite(x) AND isfinite(y)"
            % (xq, yq, qi(TABLE)))[0][0]
        return {"pairs": [{"x": r[0], "y": r[1]} for r in res], "total": total}

    def distinct(self, col, limit=500):
        self.require_data()
        cq = self.check_col(col)
        limit = max(1, min(5000, int(limit or 500)))
        res = self.q("SELECT CAST(%s AS VARCHAR) AS v, count(*) AS c FROM %s WHERE %s IS NOT NULL GROUP BY 1 ORDER BY c DESC, v LIMIT %d"
                     % (cq, qi(TABLE), cq, limit))
        return {"values": [r[0] for r in res]}

    def welch(self, g, y, a, b):
        self.require_data()
        gq, yq = self.check_col(g), self.check_col(y)
        res = self.q(
            "SELECT k, count(v), avg(v), var_samp(v) FROM (SELECT CAST(%s AS VARCHAR) AS k, TRY_CAST(%s AS DOUBLE) AS v FROM %s) "
            "WHERE k IN (?, ?) AND isfinite(v) GROUP BY k" % (gq, yq, qi(TABLE)), [str(a), str(b)])
        d = {r[0]: {"n": r[1], "mean": r[2], "var": r[3]} for r in res}
        empty = {"n": 0, "mean": None, "var": None}
        return {"a": d.get(str(a), empty), "b": d.get(str(b), empty)}

    def regression(self, x, y):
        self.require_data()
        xq, yq = self.check_col(x), self.check_col(y)
        r = self.q(
            "SELECT regr_count(yv, xv), regr_slope(yv, xv), regr_intercept(yv, xv), corr(yv, xv), regr_sxx(yv, xv), regr_syy(yv, xv) "
            "FROM (SELECT TRY_CAST(%s AS DOUBLE) AS xv, TRY_CAST(%s AS DOUBLE) AS yv FROM %s) WHERE isfinite(xv) AND isfinite(yv)"
            % (xq, yq, qi(TABLE)))[0]
        return {"n": r[0], "slope": plain(r[1]), "intercept": plain(r[2]), "r": plain(r[3]),
                "ssx": plain(r[4]), "ssy": plain(r[5])}

    # ---- changing data -----------------------------------------------------------
    def clean(self, op, **kw):
        self.require_data()
        cols = self.columns()
        if op == "drop_missing":
            before = self.row_count()
            cond = " AND ".join("%s IS NOT NULL" % qi(c) for c, _ in cols)
            self.rebuild("* EXCLUDE (__redata_rid)", "WHERE " + cond)
            return "Dropped %s rows with missing values." % format(before - self.row_count(), ",")
        if op == "fill":
            strategy = kw.get("strategy")
            marker = str(kw.get("marker") or "Unknown")
            meta = self.meta()
            missing = {i["name"]: i for i in meta["columns"]}
            select, changed = [], 0
            for c, t in cols:
                m = missing[c]["missing"]
                id_ = qi(c)
                if m and strategy == "marker":
                    if is_text_type(t):
                        select.append("COALESCE(%s, %s) AS %s" % (id_, self._lit(marker), id_))
                    else:
                        select.append("COALESCE(CAST(%s AS VARCHAR), %s) AS %s" % (id_, self._lit(marker), id_))
                    changed += m
                elif m and is_numeric_type(t) and strategy in ("zero", "mean", "median"):
                    if strategy == "zero":
                        select.append("COALESCE(%s, 0) AS %s" % (id_, id_))
                    else:
                        stat = missing[c]["mean"] if strategy == "mean" else missing[c]["median"]
                        if stat is None:
                            select.append(id_)
                            continue
                        select.append("COALESCE(CAST(%s AS DOUBLE), %r) AS %s" % (id_, float(stat), id_))
                    changed += m
                else:
                    select.append(id_)
            if changed:
                self.rebuild(", ".join(select))
            return "Filled %s missing cells. Review the changed data." % format(changed, ",")
        if op == "dedupe":
            before = self.row_count()
            part = ", ".join(qi(c) for c, _ in cols)
            self.rebuild("* EXCLUDE (__redata_rid)",
                         "QUALIFY row_number() OVER (PARTITION BY %s ORDER BY __redata_rid) = 1" % part)
            return "Removed %s exact duplicate rows." % format(before - self.row_count(), ",")
        if op == "trim":
            texts = [c for c, t in cols if is_text_type(t)]
            if not texts:
                return "Normalized 0 text cells."
            new = {c: "NULLIF(regexp_replace(trim(%s), '\\s+', ' ', 'g'), '')" % qi(c) for c in texts}
            count_expr = " + ".join("count(*) FILTER (WHERE %s IS DISTINCT FROM %s)" % (qi(c), new[c]) for c in texts)
            changes = self.q("SELECT %s FROM %s" % (count_expr, qi(TABLE)))[0][0]
            if changes:
                self.rebuild(", ".join("%s AS %s" % (new[c], qi(c)) if c in new else qi(c) for c, _ in cols))
            return "Normalized %s text cells." % format(changes, ",")
        if op == "drop_column":
            col = kw.get("column")
            self.check_col(col)
            if len(cols) <= 1:
                raise UserError("The dataset needs at least one column.")
            self.q("ALTER TABLE %s DROP COLUMN %s" % (qi(TABLE), qi(col)))
            self.changed()
            return "Removed column " + str(col)
        raise UserError("Unknown cleaning action.")

    def formula(self, name, expression):
        self.require_data()
        name = str(name or "").strip()
        if not name or not str(expression or "").strip():
            raise UserError("Enter a new column name and expression.")
        names = self.col_names()
        if name.lower() in (c.lower() for c in names):
            raise UserError("That column already exists. Choose a new name.")
        sql_expr = formula_to_sql(str(expression), names)
        self.q(
            "CREATE OR REPLACE TABLE %s AS SELECT * EXCLUDE (__redata_rid, __redata_val), "
            "CASE WHEN isfinite(__redata_val) THEN __redata_val END AS %s "
            "FROM (SELECT *, rowid AS __redata_rid, CAST(%s AS DOUBLE) AS __redata_val FROM %s) ORDER BY __redata_rid"
            % (qi(TABLE), qi(name), sql_expr, qi(TABLE)))
        self.changed()
        changed = self.q("SELECT count(%s) FROM %s" % (qi(name), qi(TABLE)))[0][0]
        return {"changed": changed, "rows": self.row_count()}

    # ---- SQL lab -------------------------------------------------------------------
    def _read_only_statement(self, query):
        query = str(query or "").strip()
        if not query:
            raise UserError("Enter a SQL query.")
        try:
            stmts = self.con.extract_statements(query)
        except duckdb.Error as e:
            raise UserError(str(e))
        if not stmts:
            raise UserError("Enter a SQL query.")
        allowed = {"StatementType.SELECT", "StatementType.EXPLAIN"}
        for s in stmts:
            if str(s.type) not in allowed:
                raise UserError(
                    "The SQL lab is read-only: only SELECT-style queries run here (DESCRIBE, SUMMARIZE and EXPLAIN "
                    "are fine). Use Clean & transform or Formula builder to change the data.")
            if _FILE_FUNC_RE.search(s.query):
                raise UserError(
                    "The SQL lab queries the loaded dataset only. File- or URL-reading functions "
                    "(read_csv, read_parquet, read_text, glob, …) are not allowed here. "
                    "Use Import data or a Connection to bring a file in.")
        return stmts[-1]

    def sql(self, query):
        stmt = self._read_only_statement(query)
        res = self.con.execute(stmt.query)
        if not res.description:
            return {"columns": [], "rows": [], "truncated": False}
        cols = unique_headers([d[0] for d in res.description], len(res.description))
        rows = res.fetchmany(SQL_UI_ROWS + 1)
        truncated = len(rows) > SQL_UI_ROWS
        rows = rows[:SQL_UI_ROWS]
        return {"columns": cols, "rows": [dict(zip(cols, [plain(v) for v in r])) for r in rows],
                "truncated": truncated}

    # ---- export -------------------------------------------------------------------------
    def export(self, kind, fmt="csv", query=None, g=None, y=None, agg="sum"):
        if kind != "sql":
            self.require_data()
        fmt = "json" if fmt == "json" else "csv"
        base = re.sub(r"\.[^.]+$", "", (self.state.get("name") or "redata").split(" · ")[0]) or "redata"
        if kind == "data":
            src = "SELECT * FROM %s ORDER BY rowid" % qi(TABLE)
            fname = "%s_export.%s" % (safe_filename(base), fmt)
        elif kind == "sql":
            stmt = self._read_only_statement(query)
            if str(stmt.type) != "StatementType.SELECT":
                raise UserError("Only SELECT results can be exported.")
            src = stmt.query.strip().rstrip(";")
            fname = "redata_sql_result.%s" % fmt
        elif kind == "aggregate":
            src = self._aggregate_export_sql(g, y, agg)
            fname = "redata_aggregate.%s" % fmt
        else:
            raise UserError("Unknown export.")
        eid = uuid.uuid4().hex
        folder = os.path.join(self.export_dir, eid)
        os.makedirs(folder)
        path = os.path.join(folder, fname)
        opts = "(FORMAT JSON, ARRAY true)" if fmt == "json" else "(HEADER, DELIMITER ',')"
        self.q("COPY (%s) TO %s %s" % (src, self._lit(path), opts))
        self.exports[eid] = path
        return {"id": eid, "file": fname, "size": os.path.getsize(path)}


# --------------------------------------------------------------------------
# HTTP layer
# --------------------------------------------------------------------------

def make_handler(engine, token, port):
    allowed_hosts = {"127.0.0.1:%d" % port, "localhost:%d" % port, "[::1]:%d" % port}

    class Handler(BaseHTTPRequestHandler):
        server_version = "ReDataEngine/" + VERSION
        protocol_version = "HTTP/1.1"

        def log_message(self, fmt, *args):
            if "/api/" in self.path:
                sys.stdout.write("  %s %s\n" % (time.strftime("%H:%M:%S"), fmt % args))

        # ---- plumbing ----
        def _send(self, code, body, ctype="application/json; charset=utf-8", extra=None):
            if isinstance(body, (dict, list)):
                body = json.dumps(body, default=str).encode("utf-8")
            elif isinstance(body, str):
                body = body.encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("Referrer-Policy", "no-referrer")
            self.send_header("X-Frame-Options", "DENY")
            self.send_header("Cross-Origin-Opener-Policy", "same-origin")
            self.send_header("Cross-Origin-Resource-Policy", "same-origin")
            self.send_header(
                "Content-Security-Policy",
                "default-src 'self'; "
                "script-src 'self' 'unsafe-inline'; "
                "style-src 'self' 'unsafe-inline' https://fonts.googleapis.com; "
                "font-src 'self' https://fonts.gstatic.com; "
                "img-src 'self' data: blob:; "
                "connect-src 'self'; "
                "frame-ancestors 'none'; base-uri 'none'; form-action 'none'")
            for k, v in (extra or {}).items():
                self.send_header(k, v)
            self.end_headers()
            if self.command != "HEAD":
                self.wfile.write(body)

        def _host_ok(self):
            return (self.headers.get("Host") or "") in allowed_hosts

        def _token_ok(self, query_token=None):
            got = self.headers.get("X-ReData-Token") or query_token or ""
            return secrets.compare_digest(str(got), token)

        def _json_body(self):
            n = int(self.headers.get("Content-Length") or 0)
            if n > 10 * 1024 * 1024:
                raise UserError("Request too large.")
            raw = self.rfile.read(n) if n else b""
            return json.loads(raw.decode("utf-8") or "{}") if raw else {}

        # ---- GET ----
        def do_GET(self):
            if not self._host_ok():
                return self._send(403, {"error": "Forbidden host."})
            url = urlparse(self.path)
            path = unquote(url.path)
            if path.startswith("/api/"):
                qs = parse_qs(url.query)
                if path.startswith("/api/download/"):
                    if not self._token_ok((qs.get("t") or [""])[0]):
                        return self._send(403, {"error": "Missing or wrong engine token."})
                    return self._download(path.rsplit("/", 1)[-1])
                if not self._token_ok():
                    return self._send(403, {"error": "Missing or wrong engine token."})
                if path == "/api/status":
                    with engine.lock:
                        return self._send(200, engine.status())
                return self._send(404, {"error": "Not found."})
            return self._static(path)

        def _static(self, path):
            if path in ("", "/"):
                path = "/index.html"
            rel = path.lstrip("/")
            parts = rel.split("/")
            if any(p.startswith(".") or p == "" for p in parts) or parts[0] in ("engine",):
                return self._send(404, "Not found", "text/plain; charset=utf-8")
            full = os.path.abspath(os.path.join(APP_ROOT, *parts))
            ext = os.path.splitext(full)[1].lower()
            if not full.startswith(APP_ROOT + os.sep) or ext not in STATIC_TYPES or not os.path.isfile(full):
                return self._send(404, "Not found", "text/plain; charset=utf-8")
            with open(full, "rb") as fh:
                body = fh.read()
            if rel == "index.html":
                inject = ('<script>window.REDATA_ENGINE=%s;</script>' %
                          json.dumps({"token": token, "version": VERSION})).encode("utf-8")
                body = re.sub(rb"<head([^>]*)>", lambda m: m.group(0) + inject, body, count=1)
            return self._send(200, body, STATIC_TYPES[ext])

        def _download(self, eid):
            path = engine.exports.get(eid)
            if not path or not os.path.isfile(path):
                return self._send(404, {"error": "This export has expired. Export again."})
            size = os.path.getsize(path)
            name = os.path.basename(path)
            ctype = "application/json" if name.endswith(".json") else "text/csv; charset=utf-8"
            self.send_response(200)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(size))
            self.send_header("Content-Disposition", 'attachment; filename="%s"' % name)
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            with open(path, "rb") as fh:
                shutil.copyfileobj(fh, self.wfile, CHUNK)

        # ---- POST ----
        def do_POST(self):
            if not self._host_ok():
                return self._send(403, {"error": "Forbidden host."})
            if not self._token_ok():
                return self._send(403, {"error": "Missing or wrong engine token."})
            url = urlparse(self.path)
            route = url.path
            started = time.time()
            try:
                if route == "/api/stage":
                    qs = parse_qs(url.query)
                    length = int(self.headers.get("Content-Length") or 0)
                    if length <= 0:
                        raise UserError("The selected file is empty.")
                    name = (qs.get("name") or ["dataset.csv"])[0]
                    # the file copy runs outside the lock; DuckDB work takes it inside
                    result = engine.stage_upload(name, self.rfile, length)
                    return self._send(200, result)
                body = self._json_body()
                if route.startswith("/api/pipeline/"):
                    result = self._route(route, body)
                else:
                    with engine.lock:
                        result = self._route(route, body)
                if isinstance(result, dict):
                    result.setdefault("ms", int((time.time() - started) * 1000))
                return self._send(200, result)
            except UserError as e:
                return self._send(400, {"error": str(e)})
            except duckdb.Error as e:
                msg = str(e).split("\n")[0] if "LINE" not in str(e) else str(e)
                return self._send(400, {"error": msg})
            except (ValueError, TypeError, KeyError, IndexError) as e:
                return self._send(400, {"error": "Bad request: %s" % e})
            except Exception as e:  # pragma: no cover
                sys.stderr.write("Unexpected error: %r\n" % e)
                return self._send(500, {"error": "Engine error: %s" % e})

        def _route(self, route, b):
            e = engine
            if route == "/api/status":
                return e.status()
            if route == "/api/stage_path":
                return e.stage_path(b.get("path"))
            if route == "/api/preview":
                return {"preview": e.excel_preview(b.get("stage"), b.get("sheet", 0))}
            if route == "/api/discard":
                e.drop_stage(b.get("stage"))
                return {"ok": True}
            if route == "/api/load":
                return {"meta": e.load(b.get("stage"), bool(b.get("header", True)), b.get("names") or None, b.get("sheet", 0))}
            if route == "/api/meta":
                return {"meta": e.meta()}
            if route == "/api/rows":
                return e.rows(b.get("offset", 0), b.get("limit", 150), b.get("sort"), b.get("dir", 1),
                              b.get("column"), b.get("query", ""))
            if route == "/api/aggregate":
                return e.aggregate(b.get("g"), b.get("y"), b.get("agg", "sum"), b.get("order", "desc"),
                                   b.get("limit", 20), bool(b.get("groups")))
            if route == "/api/histogram":
                return e.histogram(b.get("col"), b.get("bins", 10))
            if route == "/api/box":
                return e.box(b.get("col"))
            if route == "/api/scatter":
                return e.scatter(b.get("x"), b.get("y"), b.get("limit", 5000))
            if route == "/api/distinct":
                return e.distinct(b.get("col"), b.get("limit", 500))
            if route == "/api/welch":
                return e.welch(b.get("g"), b.get("y"), b.get("a"), b.get("b"))
            if route == "/api/regression":
                return e.regression(b.get("x"), b.get("y"))
            if route == "/api/clean":
                msg = e.clean(b.get("op"), **{k: v for k, v in b.items() if k != "op"})
                return {"message": msg, "meta": e.meta()}
            if route == "/api/formula":
                r = e.formula(b.get("name"), b.get("expr"))
                r["meta"] = e.meta()
                return r
            if route == "/api/sql":
                return e.sql(b.get("query"))
            if route.startswith("/api/pipeline/"):
                return self._pipeline_route(route[len("/api/pipeline/"):], b)
            if route == "/api/open_staged":
                return {"meta": e.open_staged(b.get("table"), bool(b.get("link")))}
            if route == "/api/export":
                return e.export(b.get("kind"), b.get("format", "csv"), b.get("query"), b.get("g"), b.get("y"),
                                b.get("agg", "sum"))
            raise UserError("Unknown engine endpoint.")

        def _pipeline_route(self, name, b):
            from pipeline import PipelineError
            p = engine.pipeline
            if p is None:
                raise UserError("The pipeline is not running.")
            try:
                if name == "catalog":
                    return {"catalog": p.catalog()}
                if name == "overview":
                    return p.overview()
                if name == "save_connection":
                    return p.save_connection(b)
                if name == "delete_connection":
                    return p.delete_connection(b.get("id"))
                if name == "test_connection":
                    return p.test_connection(b.get("id"))
                if name == "snowflake_key":
                    return p.snowflake_key(b.get("id"))
                if name == "source_tables":
                    return p.source_tables(b.get("id"))
                if name == "source_columns":
                    return p.source_columns(b.get("id"), b.get("source"))
                if name == "save_stream":
                    return p.save_stream(b)
                if name == "delete_stream":
                    return p.delete_stream(b.get("id"), bool(b.get("drop")))
                if name == "reset_stream":
                    return p.reset_stream(b.get("id"))
                if name == "sync_now":
                    return p.enqueue(b.get("id"))
                if name == "runs":
                    return {"runs": p.runs(int(b.get("limit") or 50), b.get("stream"))}
            except PipelineError as err:
                raise UserError(str(err))
            raise UserError("Unknown pipeline action.")

    return Handler


def main():
    ap = argparse.ArgumentParser(description="ReData local engine (DuckDB).")
    ap.add_argument("--port", type=int, default=8765)
    ap.add_argument("--workspace", default=os.path.join(ENGINE_DIR, "workspace"),
                    help="folder for the engine database, uploads and exports")
    ap.add_argument("--memory-limit", default=None, help="DuckDB memory limit, e.g. 4GB (default: 80%% of RAM)")
    ap.add_argument("--threads", type=int, default=None, help="DuckDB worker threads (default: all cores)")
    ap.add_argument("--no-browser", action="store_true", help="do not open the browser automatically")
    args = ap.parse_args()

    if not os.path.isfile(os.path.join(APP_ROOT, "index.html")):
        sys.exit("index.html was not found next to the engine folder (%s)." % APP_ROOT)
    engine = Engine(os.path.abspath(args.workspace), args.memory_limit, args.threads)
    engine.start_pipeline()
    token = secrets.token_urlsafe(24)

    server, port = None, args.port
    for candidate in range(args.port, args.port + 20):
        try:
            server = ThreadingHTTPServer(("127.0.0.1", candidate), make_handler(engine, token, candidate))
            port = candidate
            break
        except OSError:
            continue
    if server is None:
        sys.exit("No free port between %d and %d." % (args.port, args.port + 19))
    server.daemon_threads = True
    url = "http://127.0.0.1:%d/" % port
    print("")
    print("  ReData local engine %s  (DuckDB %s)" % (VERSION, duckdb.__version__))
    print("  Open:       %s" % url)
    print("  Workspace:  %s" % engine.workspace)
    if engine.state.get("name"):
        print("  Dataset:    %s (restored from last session)" % engine.state["name"])
    if engine.pipeline and engine.pipeline.cfg["streams"]:
        n = len(engine.pipeline.cfg["streams"])
        print("  Syncs:      %d configured; scheduled ones run while this window is open" % n)
    print("  Stop:       press Ctrl+C in this window")
    print("")
    if not args.no_browser:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n  Stopping ReData engine…")
    finally:
        server.server_close()
        with engine.lock:
            if engine.pipeline is not None:
                engine.pipeline.stop.set()
                for cur in (getattr(engine.pipeline, "cur", None), getattr(engine.pipeline, "rcur", None)):
                    try:
                        if cur is not None:
                            cur.close()
                    except Exception:
                        pass
            engine.con.close()


if __name__ == "__main__":
    main()
